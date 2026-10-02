"""RT2 Part 2 (RT2_RULING_1 R1) — build Burt's CORRECTED workbook set from his sealed 21 Sep attachment.

    python build_corrected_set.py <nuwepryslyste.zip> <out-dir>

Writes, OUTSIDE git (Burt's files never enter the repo, and the original is never edited):

    <out>/original/    the three zip members, byte for byte
    <out>/c1/          correction 1 only (PU!C17 4095 -> 4100, in both places below); FORMULAS = original
    <out>/corrected/   corrections 1 + 2 — the set the golden is built on
    <out>/corrections.json   provenance: the zip, every member, every corrected file (sha256), and
                             each correction's cells, from, to and authority

The two corrections, both authorised by Burt in writing and neither in his file yet:

  1. PU!C17  4095 -> 4100. Burt, 1 Oct 2026: 32D PU FOAM is R4 100 (the R4 095 on his list was a
     mistake). Applied in BOTH places the value lives:
       - PRICE 2017 MARCH.xlsx PU!C17 — a literal (<v>4095</v>); no formula in PRICE references it;
       - the GRP's own cache of it: xl/externalLinks/externalLink1.xml (target PRICE 2017 MARCH.xlsx),
         sheet PU, cell C17. LibreOffice headless prices every [1] reference from THIS cache and never
         reads the adjacent PRICE file — proven (RT2_RETURN_2 §2): PRICE PU!C17 = 9999 moved no
         section of a 12-section PU body. So the cache is what the golden actually sees.
     Each is one token in one XML part.
  2. GRP Costings 2018.xlsx, every GALV PLATE line inside a SUB FRAME + LIGHT BOX ASSY block whose
     price reads MILD STEEL $D$28 (1.2MM): repoint the reference to $D$27 (1.0MM). The label "1.2MM"
     is left exactly as Burt wrote it (the 17 Sep "price re-pointed, description not updated"
     pattern). Burt's email, 7 Sep 2026 15:47 (to Michael and Nadie): "Sorry that should be 1,0mm
     Galv Plate. We have changed it but never corrected the costing."  A GALV PLATE line whose
     price reads anything other than $D$28 is REPORTED and left (RT2_RETURN_2 §2).

Every touched cell is guarded: the patch refuses unless the cell holds EXACTLY the expected
original text. Members that are not touched are copied with their content unchanged (re-zipped,
so their compressed bytes may differ — the proof is on content: diff_cells() below).
"""
from __future__ import annotations

import hashlib
import io
import json
import re
import sys
import zipfile
from pathlib import Path

GRP, PRICE, FORMULAS = "GRP Costings 2018.xlsx", "PRICE 2017 MARCH.xlsx", "FORMULAS 2018.xls"
ZIP_SHA = "ee8e2159031be94738797b72f0b9a11a7baa1ee5eabd3a4f34269606106eea7b"
MEMBER_SHA = {
    GRP: "c0926043ff87a1d117c65d6290ce9c144daf58f0df871ed6e0c11a5baf9547db",
    PRICE: "073cdc8409f33f4259737f1cf2961530898116121de31f9c5133a404ea0945a5",
    FORMULAS: "f90db787b47a2f81763a5518fa2e0c1ddbbb4c502b2f828861283aa50e1e2f3f",
}

C1 = {"file": PRICE, "sheet": "PU", "cell": "C17", "from": "4095", "to": "4100",
      "what": "32D PU FOAM sheet price",
      "authority": "Burt, 1 Oct 2026 (relayed by Michael, RT2_RULING_1 R1): 32D PU FOAM is R4 100"}
C2_FROM, C2_TO = "'[1]MILD STEEL'!$D$28/2.98", "'[1]MILD STEEL'!$D$27/2.98"
C2_CELLS = [  # (sheet, price cell) — found by the audit tool's own discovery (RT2_RETURN_2 §2)
    ("UP TO 2.3 CHILLER BODY", "G163"),
    ("UP TO 5.5 CHILLER AND 2.3 WIDE", "G187"),
    (" 4.9 & UP CHILLER AND 2.5 WIDE ", "R185"),
    (" UP TO 4.8 MT FREEZER  (2", "G196"),
    ("EXPLOSIVE UP TO 2.7", "G165"),
    ("EXPLOSIVE 2.7 TO 4.8", "G175"),
    ("EXPLOSIVE 4.9 AND UP", "G175"),
]
C2 = {"file": GRP, "section": "SUB FRAME + LIGHT BOX ASSY", "line": "1.2MM GALV PLATE (label kept)",
      "from": "=" + C2_FROM, "to": "=" + C2_TO, "cells": [f"{s.strip()}!{c}" for s, c in C2_CELLS],
      "authority": "Burt's email of 7 Sep 2026 15:47 to Michael and Nadie: \"Sorry that should be 1,0mm "
                   "Galv Plate. We have changed it but never corrected the costing.\" Michael, same thread: "
                   "\"...the case on all 'SUB FRAME + LIGHT BOX ASSY' where the 1.2MM GALV PLATE is being used.\""}


def sha(b: bytes) -> str:
    return hashlib.sha256(b).hexdigest()


def sheet_parts(xlsx: bytes) -> dict[str, str]:
    """{sheet name: 'xl/worksheets/sheetN.xml'} from workbook.xml + its rels."""
    with zipfile.ZipFile(io.BytesIO(xlsx)) as z:
        wb = z.read("xl/workbook.xml").decode("utf-8")
        rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    rel_target = dict(re.findall(r'<Relationship [^>]*Id="([^"]+)"[^>]*Target="([^"]+)"', rels))
    rel_target.update({i: t for t, i in re.findall(r'<Relationship [^>]*Target="([^"]+)"[^>]*Id="([^"]+)"', rels)})
    out = {}
    for name, rid in re.findall(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', wb):
        name = (name.replace("&amp;", "&").replace("&apos;", "'").replace("&quot;", '"')
                .replace("&lt;", "<").replace("&gt;", ">"))
        t = rel_target[rid].lstrip("/")
        out[name] = t if t.startswith("xl/") else "xl/" + t
    return out


def rezip(xlsx: bytes, replace: dict[str, bytes]) -> bytes:
    """The same members in the same order with the same compression, `replace` swapped in."""
    src = zipfile.ZipFile(io.BytesIO(xlsx))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as dst:
        for info in src.infolist():
            data = replace.get(info.filename, src.read(info.filename))
            zi = zipfile.ZipInfo(info.filename, date_time=info.date_time)
            zi.compress_type, zi.external_attr = info.compress_type, info.external_attr
            dst.writestr(zi, data)
    missing = set(replace) - {i.filename for i in src.infolist()}
    if missing:
        raise SystemExit(f"REFUSED: no such member(s) {sorted(missing)}")
    return buf.getvalue()


def patch_c1(price: bytes) -> bytes:
    part = sheet_parts(price)[C1["sheet"]]
    xml = zipfile.ZipFile(io.BytesIO(price)).read(part).decode("utf-8")
    old = f'<c r="{C1["cell"]}"><v>{C1["from"]}</v></c>'
    if xml.count(old) != 1:
        raise SystemExit(f"REFUSED: {C1['sheet']}!{C1['cell']} is not exactly {old!r} in {part}")
    if re.search(r"<f>[^<]*\bC\$?17\b", xml):
        raise SystemExit("REFUSED: a PU formula references C17 — the literal is not alone")
    new = xml.replace(old, f'<c r="{C1["cell"]}"><v>{C1["to"]}</v></c>')
    return rezip(price, {part: new.encode("utf-8")})


def price_link_part(grp: bytes) -> str:
    """The externalLink part whose target is PRICE 2017 MARCH.xlsx."""
    z = zipfile.ZipFile(io.BytesIO(grp))
    hits = []
    for n in z.namelist():
        m = re.match(r"xl/externalLinks/_rels/(externalLink\d+\.xml)\.rels$", n)
        if m and re.search(r'Target="[^"]*PRICE%202017%20MARCH\.xlsx"', z.read(n).decode("utf-8")):
            hits.append("xl/externalLinks/" + m.group(1))
    if len(hits) != 1:
        raise SystemExit(f"REFUSED: expected exactly one external link to {PRICE}, found {hits}")
    return hits[0]


def patch_c1_cache(grp: bytes) -> tuple[bytes, str]:
    """The GRP's cached [1]PU!C17 (what LibreOffice prices from) 4095 -> 4100."""
    part = price_link_part(grp)
    xml = zipfile.ZipFile(io.BytesIO(grp)).read(part).decode("utf-8")
    names = re.findall(r'<sheetName val="([^"]+)"', xml)
    sid = names.index(C1["sheet"])
    m = re.search(rf'<sheetData sheetId="{sid}"[^>]*>(.*?)</sheetData>', xml, re.S)
    if not m:
        raise SystemExit(f"REFUSED: {part} caches nothing for sheet {C1['sheet']!r}")
    body = m.group(1)
    old = re.findall(rf'<cell r="{C1["cell"]}"(?: [^>]*)?><v>{C1["from"]}</v></cell>', body)
    if len(old) != 1:
        raise SystemExit(f"REFUSED: the cached {C1['sheet']}!{C1['cell']} in {part} is not exactly {C1['from']}")
    new_body = body.replace(old[0], old[0].replace(f"<v>{C1['from']}</v>", f"<v>{C1['to']}</v>"))
    return rezip(grp, {part: (xml[:m.start(1)] + new_body + xml[m.end(1):]).encode("utf-8")}), part


def diff_link_cache(a: bytes, b: bytes) -> list[str]:
    """Every cached external value that differs between two GRP files (all links, all sheets)."""
    def cache(x: bytes) -> dict:
        z = zipfile.ZipFile(io.BytesIO(x))
        out = {}
        for n in sorted(z.namelist()):
            if not re.match(r"xl/externalLinks/externalLink\d+\.xml$", n):
                continue
            xml = z.read(n).decode("utf-8")
            names = re.findall(r'<sheetName val="([^"]+)"', xml)
            for sid, body in re.findall(r'<sheetData sheetId="(\d+)"[^>/]*>(.*?)</sheetData>', xml, re.S):
                for r, v in re.findall(r'<cell r="([A-Z]+[0-9]+)"[^>]*>\s*<v>([^<]*)</v>', body):
                    out[f"{n.rsplit('/', 1)[-1]} {names[int(sid)]}!{r}"] = v
        return out
    ca, cb = cache(a), cache(b)
    return [f"cache {k}: {ca.get(k)!r} -> {cb.get(k)!r}" for k in sorted(set(ca) | set(cb)) if ca.get(k) != cb.get(k)]


def patch_c2(grp: bytes) -> tuple[bytes, list[dict]]:
    parts = sheet_parts(grp)
    z = zipfile.ZipFile(io.BytesIO(grp))
    by_part: dict[str, str] = {}
    done = []
    for sheet, cell in C2_CELLS:
        part = parts[sheet]
        xml = by_part.get(part) or z.read(part).decode("utf-8")
        m = re.search(rf'<c r="{cell}"[^>]*>(.*?)</c>', xml, re.S)
        if not m:
            raise SystemExit(f"REFUSED: no cell {sheet!r}!{cell}")
        body = m.group(1)
        f = re.search(r"<f(?: [^>]*)?>([^<]*)</f>", body)
        if not f or f.group(1).replace("&apos;", "'") != C2_FROM or "t=\"shared\"" in body:
            raise SystemExit(f"REFUSED: {sheet!r}!{cell} formula is {f.group(1) if f else None!r}, expected {C2_FROM!r}")
        raw_from = f.group(1)
        raw_to = raw_from.replace("$D$28", "$D$27")
        new_body = body.replace(f">{raw_from}</f>", f">{raw_to}</f>", 1)
        xml = xml[:m.start(1)] + new_body + xml[m.end(1):]
        by_part[part] = xml
        done.append({"sheet": sheet, "cell": cell, "part": part, "from": "=" + C2_FROM, "to": "=" + C2_TO})
    return rezip(grp, {p: x.encode("utf-8") for p, x in by_part.items()}), done


def diff_cells(a: bytes, b: bytes) -> list[str]:
    """Every cell whose FORMULA text or VALUE differs between two xlsx files (all sheets)."""
    from openpyxl import load_workbook
    out = []
    for data_only in (False, True):
        wa = load_workbook(io.BytesIO(a), data_only=data_only, read_only=True)
        wb = load_workbook(io.BytesIO(b), data_only=data_only, read_only=True)
        for ws in wa.worksheets:
            ca = {c.coordinate: c.value for row in ws.iter_rows() for c in row if hasattr(c, "coordinate") and c.value is not None}
            cb = {c.coordinate: c.value for row in wb[ws.title].iter_rows() for c in row if hasattr(c, "coordinate") and c.value is not None}
            for k in sorted(set(ca) | set(cb)):
                if ca.get(k) != cb.get(k):
                    out.append(f"{'value' if data_only else 'formula'} {ws.title.strip()}!{k}: {ca.get(k)!r} -> {cb.get(k)!r}")
    return out


def main(zip_path: str, out_dir: str) -> None:
    zb = Path(zip_path).read_bytes()
    if sha(zb) != ZIP_SHA:
        raise SystemExit(f"REFUSED: {zip_path} sha256 {sha(zb)} is not Burt's confirmed 21 Sep attachment")
    members = {Path(n).name: zipfile.ZipFile(io.BytesIO(zb)).read(n) for n in zipfile.ZipFile(io.BytesIO(zb)).namelist()}
    for n, want in MEMBER_SHA.items():
        if sha(members[n]) != want:
            raise SystemExit(f"REFUSED: zip member {n} sha256 {sha(members[n])} != {want}")
    out = Path(out_dir)
    sets = {"original": dict(members)}
    price1 = patch_c1(members[PRICE])
    grp1, cache_part = patch_c1_cache(members[GRP])
    grp2, c2_done = patch_c2(grp1)
    sets["c1"] = {GRP: grp1, PRICE: price1, FORMULAS: members[FORMULAS]}
    sets["corrected"] = {GRP: grp2, PRICE: price1, FORMULAS: members[FORMULAS]}
    for name, files in sets.items():
        d = out / name
        d.mkdir(parents=True, exist_ok=True)
        for fn, data in files.items():
            (d / fn).write_bytes(data)

    d1 = diff_cells(members[PRICE], price1)
    d1g = diff_cells(members[GRP], grp1) + diff_link_cache(members[GRP], grp1)
    d2 = diff_cells(grp1, grp2) + diff_link_cache(grp1, grp2)
    print("== correction 1 — PRICE, every cell that differs from the original:")
    print("\n".join(f"   {x}" for x in d1))
    print("== correction 1 — GRP, every worksheet cell AND every cached external value that differs:")
    print("\n".join(f"   {x}" for x in d1g))
    print("== correction 2 — GRP (over correction 1), every worksheet cell and cached external value that differs:")
    print("\n".join(f"   {x}" for x in d2))
    want1 = {f"formula PU!C17: 4095 -> 4100", f"value PU!C17: 4095 -> 4100"}
    if set(d1) != want1:
        raise SystemExit(f"REFUSED: correction 1 changed more (or less) than PU!C17: {d1}")
    want1g = {f"cache {cache_part.rsplit('/', 1)[-1]} PU!C17: '4095' -> '4100'"}
    if set(d1g) != want1g:
        raise SystemExit(f"REFUSED: correction 1 changed more (or less) than the GRP's cached PU!C17: {d1g}")
    want2 = {f"formula {s.strip()}!{c}: {'=' + C2_FROM!r} -> {'=' + C2_TO!r}" for s, c in C2_CELLS}
    if set(d2) != want2:
        raise SystemExit(f"REFUSED: correction 2 changed more (or less) than the {len(C2_CELLS)} cells: {d2}")

    prov = {
        "source": {"zip": str(Path(zip_path).name), "zip_sha256": ZIP_SHA, "received": "21 Sep 2026 08:39 (Burt's email)",
                   "confirmed": "Michael, 1 Oct 2026 (RT2_RULING_1 R1)", "members": MEMBER_SHA},
        "corrections": [dict(C1, id=1, cells_changed=d1 + d1g, applied_to=[f"{PRICE} PU!C17",
                                                                          f"{GRP} {cache_part} cached [1]PU!C17"]),
                        dict(C2, id=2, cells_changed=d2, patched=c2_done)],
        "note": "LibreOffice headless prices external references from the GRP's own link cache, never from the "
                "adjacent PRICE file (PRICE PU!C17 = 9999 moved nothing) — so a price correction is made in the "
                "cache too, and a new price list only reaches the golden through a GRP saved with its links updated.",
        "files": {name: {fn: sha(data) for fn, data in files.items()} for name, files in sets.items()},
        "first_check_on_burts_next_set": "both corrections must already be in his file (RT2_RULING_1 R1)",
    }
    (out / "corrections.json").write_text(json.dumps(prov, indent=1), encoding="utf-8")
    print(f"== wrote {out}/original, {out}/c1, {out}/corrected and corrections.json")
    for name, files in sets.items():
        print(f"   {name:<10} " + " · ".join(f"{fn.split()[0]} {sha(data)[:12]}" for fn, data in files.items()))


if __name__ == "__main__":
    main(*sys.argv[1:])
