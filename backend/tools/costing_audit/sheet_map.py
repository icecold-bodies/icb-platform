"""Auto-discovery of a costing sheet's cell map (ratified default 5).

Burt's body sheets share one shape, anchored on a LABEL column (A on most
sheets, L on the live block of ` 4.9 & UP CHILLER AND 2.5 WIDE `):

    row 4-6   LENGTH / WIDTH / HEIGHT   label at L, value at L+2
    row 8..   flag block                label at L, thickness at L+2, Y/N at L+3
              (FRONT EPS, FRONT PU, DRD EPS, ... then RICE GRAIN FLOOR etc.)
    then, per section:
              label row                 text at L (or L+1), nothing else in L..L+9
              header row                a row holding both PRICE and TOTAL
              line rows                 description at L, qty at PRICE-1, PRICE, TOTAL
              TOTAL row                 raw =SUM(..) at the TOTAL column, an optional
                                        gate =IF(Dnn="Y",1,0) at L+8 and the GATED
                                        section total at L+9 (the J column)
    last      GRAND TOTAL               =SUM(J..) — the sum of every gated J cell

A section's material cost is the sum of its gated J cells (exactly what the
GRAND TOTAL formula adds), which also absorbs Burt's variants: a second
PU-gated REAR FRAME row, SPRAY PAINTING with no header/TOTAL row, REFLEXITE
TAPE whose single line carries J directly.

Everything discovered is written out by `audit discover` for a human
tick-through; `sheet_maps/<slug>.yaml` overrides win over discovery.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field, asdict
from pathlib import Path

from openpyxl.utils import get_column_letter, column_index_from_string

from .mapping import norm_name, DOOR_PANELS

HEADER_WORDS = {"WIDTH", "HEIGHT", "LENGTH", "M2", "PRICE", "TOTAL", "QUANT", "QUANTI",
                "QUANTITY", "TOTAL M", "TOTAL MTR", "QTY"}
INPUT_LABELS = ("LENGTH", "WIDTH", "HEIGHT")
BLOCK_WIDTH = 10          # label col .. gated col (A..J)
GATE_OFFSET = 8           # I
GATED_OFFSET = 9          # J
FLAG_FIRST_ROW = 8
FLAG_ROW_LIMIT = 40
_GATE_REF_RE = re.compile(r'\$?([A-Z]{1,2})\$?(\d+)\s*=\s*"Y"', re.I)
_EXT_REF_RE = re.compile(r"\[(\d+)\]")


def _is_formula(v) -> bool:
    return isinstance(v, str) and v.startswith("=")


def _text(v) -> str:
    return v.strip() if isinstance(v, str) and not v.startswith("=") else ""


@dataclass
class FlagRow:
    label: str
    row: int
    thickness_cell: str
    flag_cell: str
    thickness: float | None
    flag: str | None            # cached Y/N (after Burt's own formulas, if any)
    is_formula: bool            # True when the sheet drives the cell by formula


@dataclass
class LineRow:
    row: int
    desc: str
    qty_cell: str | None
    price_cell: str | None
    total_cell: str | None
    stale_link: bool = False


@dataclass
class GateCell:
    cell: str
    formula: str
    flag_cells: list[str]
    flag_labels: list[str]


@dataclass
class Section:
    name: str                   # display name (door-fittings prefixed, de-duplicated)
    label: str                  # the raw label text on the sheet
    label_row: int
    first_row: int
    last_row: int
    header_row: int | None
    total_row: int | None
    qty_col: str | None
    price_col: str | None
    total_col: str | None
    raw_total_cell: str | None
    gated_cells: list[str]
    gate_cells: list[GateCell]
    lines: list[LineRow]
    cached_total: float | None      # sum of gated cells, from the saved workbook
    unverifiable: str | None = None # e.g. STALE_LINK


@dataclass
class SheetMap:
    sheet: str
    label_col: str
    inputs: dict[str, str]                 # LENGTH -> "C4"
    input_values: dict[str, float | None]  # LENGTH -> 5.5 (saved value)
    flags: list[FlagRow]
    sections: list[Section]
    grand_total_cell: str | None
    grand_total_cached: float | None
    selfcheck_ok: bool | None
    selfcheck_delta: float | None
    stale_links: dict[int, str]            # external-link index -> missing target
    control_panel: bool                    # formula-driven flag block (UP TO 5.5 CHILLER)
    warnings: list[str] = field(default_factory=list)
    source: str = "auto"

    # -- lookups -----------------------------------------------------------
    def flag_by_label(self, label: str) -> FlagRow | None:
        want = norm_name(label)
        for f in self.flags:
            if norm_name(f.label) == want:
                return f
        return None

    def flag_by_cell(self, cell: str) -> FlagRow | None:
        for f in self.flags:
            if f.flag_cell == cell:
                return f
        return None

    def to_dict(self) -> dict:
        return asdict(self)


# ── helpers ────────────────────────────────────────────────────────────────

def _cell(col: int, row: int) -> str:
    return f"{get_column_letter(col)}{row}"


def _find_label_cols(ws) -> list[int]:
    """Columns whose row-4 cell reads LENGTH — one per input block."""
    cols = []
    for c in range(1, ws.max_column + 1):
        if _text(ws.cell(4, c).value).upper() == "LENGTH":
            cols.append(c)
    return cols


def stale_external_links(wb, workbook_dir: Path | None) -> dict[int, str]:
    """External links whose target workbook is not adjacent. Formulas carry
    them as [k]; every cell that references a stale k is unrecalculable."""
    out: dict[int, str] = {}
    links = getattr(wb, "_external_links", None) or []
    for i, link in enumerate(links, start=1):
        target = ""
        try:
            target = link.file_link.Target or ""
        except AttributeError:
            pass
        name = Path(target.replace("%20", " ")).name
        if workbook_dir is None or not (Path(workbook_dir) / name).is_file():
            out[i] = name or f"link {i}"
    return out


def _formula_is_stale(formula, stale: dict[int, str]) -> bool:
    if not stale or not _is_formula(formula):
        return False
    return any(int(k) in stale for k in _EXT_REF_RE.findall(formula))


# ── discovery ──────────────────────────────────────────────────────────────

def discover_sheet(ws, wsv, *, override: dict | None = None,
                   stale_links: dict[int, str] | None = None) -> SheetMap:
    """Build the SheetMap for one sheet. `ws` = formulas view, `wsv` = the
    same sheet from a data_only load (cached values). `override` is the
    sheet_maps/*.yaml document for this sheet, if any."""
    override = override or {}
    stale_links = stale_links or {}
    warnings: list[str] = []

    label_cols = _find_label_cols(ws)
    if override.get("label_col"):
        L = column_index_from_string(str(override["label_col"]).upper())
    elif label_cols:
        L = label_cols[0]
        if len(label_cols) > 1:
            # Prefer a block whose GRAND TOTAL recalculates (the dead left block of
            # the 4.9 chiller is #REF! from a missing 2004 price file).
            for cand in label_cols:
                gt = _find_grand_total(ws, wsv, cand)
                if gt is not None and isinstance(gt[1], (int, float)):
                    L = cand
                    break
            warnings.append("several LENGTH blocks at columns %s; using %s"
                            % (", ".join(get_column_letter(c) for c in label_cols), get_column_letter(L)))
    else:
        raise ValueError("no LENGTH label in row 4 — not a standard body sheet (add a sheet_map)")

    inputs: dict[str, str] = {}
    input_values: dict[str, float | None] = {}
    for r, want in zip((4, 5, 6), INPUT_LABELS):
        got = _text(ws.cell(r, L).value).upper()
        if got != want:
            warnings.append(f"row {r}: expected {want}, found {got!r}")
        inputs[want] = _cell(L + 2, r)
        v = wsv.cell(r, L + 2).value
        input_values[want] = float(v) if isinstance(v, (int, float)) else None

    # flag block: contiguous labelled rows from row 8
    flags: list[FlagRow] = []
    r = FLAG_FIRST_ROW
    control_panel = False
    while r < FLAG_ROW_LIMIT:
        label = _text(ws.cell(r, L).value)
        if not label:
            break
        tv, fv = ws.cell(r, L + 2).value, ws.cell(r, L + 3).value
        is_f = _is_formula(tv) or _is_formula(fv)
        control_panel = control_panel or is_f
        tcached = wsv.cell(r, L + 2).value
        fcached = wsv.cell(r, L + 3).value
        flags.append(FlagRow(
            label=label, row=r,
            thickness_cell=_cell(L + 2, r), flag_cell=_cell(L + 3, r),
            thickness=float(tcached) if isinstance(tcached, (int, float)) else None,
            flag=(str(fcached).strip().upper() if fcached is not None else None),
            is_formula=is_f))
        r += 1
    last_flag_row = flags[-1].row if flags else FLAG_FIRST_ROW
    if not flags:
        warnings.append("no flag block found at row 8")

    gt = _find_grand_total(ws, wsv, L)
    grand_total_cell = gt[0] if gt else None
    grand_total_cached = gt[1] if gt and isinstance(gt[1], (int, float)) else None
    grand_total_row = int(re.sub(r"[A-Z]+", "", grand_total_cell)) if grand_total_cell else ws.max_row + 1
    if gt and grand_total_cached is None:
        warnings.append(f"GRAND TOTAL {grand_total_cell} is not numeric: {gt[1]!r}")

    # classify rows after the flag block
    label_rows: list[tuple[int, str]] = []
    header_rows: dict[int, tuple[int, int]] = {}      # row -> (price_col, total_col)
    total_rows: set[int] = set()
    for r in range(last_flag_row + 1, grand_total_row):
        vals = {c: ws.cell(r, c).value for c in range(L, L + BLOCK_WIDTH)}
        texts = {c: _text(v).upper() for c, v in vals.items()}
        price_cols = [c for c, t in texts.items() if t.startswith("PRICE")]
        total_cols = [c for c, t in texts.items() if t in ("TOTAL", "TOTAL M", "TOTAL MTR")]
        if price_cols and total_cols and any(c > price_cols[0] for c in total_cols):
            header_rows[r] = (price_cols[0], max(c for c in total_cols if c > price_cols[0]))
            continue
        if any(t == "TOTAL" for t in texts.values()):
            total_rows.add(r)
            continue
        lab = _text(vals[L]) or _text(vals[L + 1])
        others = [v for c, v in vals.items() if c not in (L, L + 1) and v is not None]
        if lab and not others and lab.upper() not in HEADER_WORDS and not _is_formula(vals[L]):
            # a bare text in the label column(s) with nothing else in the block
            label_rows.append((r, lab))

    # sections
    sections: list[Section] = []
    last_door: str | None = None
    seen_names: dict[str, int] = {}
    bounds = label_rows + [(grand_total_row, None)]
    for (lr, lab), (nr, _) in zip(bounds, bounds[1:]):
        first, last = lr + 1, nr - 1
        hdr = next((h for h in sorted(header_rows) if first <= h <= last), None)
        tot = next((t for t in sorted(total_rows) if (hdr or first) < t <= last), None)
        qty_col = price_col = total_col = None
        if hdr is not None:
            pc, tc = header_rows[hdr]
            price_col, total_col, qty_col = get_column_letter(pc), get_column_letter(tc), get_column_letter(pc - 1)
        # lines: rows with a description between header (or label) and TOTAL (or end)
        lo = (hdr + 1) if hdr is not None else first
        hi = (tot - 1) if tot is not None else last
        lines: list[LineRow] = []
        for r in range(lo, hi + 1):
            desc = _text(ws.cell(r, L).value)
            if not desc or desc.upper() in HEADER_WORDS:
                continue
            row_vals = [ws.cell(r, c).value for c in range(L + 1, L + BLOCK_WIDTH)]
            if all(v is None for v in row_vals):
                continue
            stale = any(_formula_is_stale(v, stale_links) for v in row_vals)
            lines.append(LineRow(
                row=r, desc=desc,
                qty_cell=f"{qty_col}{r}" if qty_col else None,
                price_cell=f"{price_col}{r}" if price_col else None,
                total_cell=f"{total_col}{r}" if total_col else None,
                stale_link=stale))
        gated_cells = [_cell(L + GATED_OFFSET, r) for r in range(first, last + 1)
                       if ws.cell(r, L + GATED_OFFSET).value is not None]
        gates: list[GateCell] = []
        for r in range(first, last + 1):
            v = ws.cell(r, L + GATE_OFFSET).value
            if _is_formula(v) and "IF(" in v.upper():
                refs = [f"{c.upper()}{n}" for c, n in _GATE_REF_RE.findall(v)]
                labels = []
                for ref in refs:
                    fr = next((f for f in flags if f.flag_cell == ref), None)
                    labels.append(fr.label if fr else "?")
                gates.append(GateCell(cell=_cell(L + GATE_OFFSET, r), formula=v,
                                      flag_cells=refs, flag_labels=labels))
        # name: bare DOOR FITTINGS takes the enclosing door section's prefix
        name = lab
        up = norm_name(lab)
        if up in DOOR_PANELS:
            last_door = up
        elif up == "DOOR FITTINGS" and last_door:
            name = f"{last_door} DOOR FITTINGS"
        n = seen_names.get(norm_name(name), 0) + 1
        seen_names[norm_name(name)] = n
        if n > 1:
            name = f"{name} ({n})"
        cached = 0.0
        numeric = True
        for cell in gated_cells:
            v = wsv[cell].value
            if isinstance(v, (int, float)):
                cached += float(v)
            elif v is not None:
                numeric = False
        unverifiable = None
        if any(l.stale_link for l in lines):
            unverifiable = "STALE_LINK"
        elif not numeric:
            unverifiable = "ERROR_CELL"
        sections.append(Section(
            name=name, label=lab, label_row=lr, first_row=first, last_row=last,
            header_row=hdr, total_row=tot, qty_col=qty_col, price_col=price_col, total_col=total_col,
            raw_total_cell=f"{total_col}{tot}" if (tot is not None and total_col) else None,
            gated_cells=gated_cells, gate_cells=gates, lines=lines,
            cached_total=round(cached, 6) if numeric else None, unverifiable=unverifiable))

    delta = ok = None
    if grand_total_cached is not None:
        total = sum(s.cached_total or 0.0 for s in sections)
        delta = round(total - grand_total_cached, 4)
        ok = abs(delta) < 0.01
        if not ok:
            warnings.append(f"self-check: section sum {total:.2f} != GRAND TOTAL {grand_total_cached:.2f} (delta {delta:+.2f})")
    if not sections:
        warnings.append("no sections discovered")

    sm = SheetMap(sheet=ws.title, label_col=get_column_letter(L), inputs=inputs,
                  input_values=input_values, flags=flags,
                  sections=sections, grand_total_cell=grand_total_cell,
                  grand_total_cached=grand_total_cached, selfcheck_ok=ok, selfcheck_delta=delta,
                  stale_links=dict(stale_links), control_panel=control_panel, warnings=warnings,
                  source="override" if override else "auto")
    _apply_section_overrides(sm, override)
    return sm


def _find_grand_total(ws, wsv, L: int):
    """(cell, cached value) of the GRAND TOTAL for the block anchored at L."""
    for r in range(ws.max_row, 6, -1):
        for c in range(L, L + BLOCK_WIDTH):
            if _text(ws.cell(r, c).value).upper() == "GRAND TOTAL":
                cell = _cell(L + GATED_OFFSET, r)
                return cell, wsv[cell].value
    return None


def _apply_section_overrides(sm: SheetMap, override: dict) -> None:
    """sheet_maps/*.yaml may rename sections (`rename: {"OLD": "NEW"}`) or drop
    them (`ignore: [..]`) after discovery."""
    ren = {norm_name(k): v for k, v in (override.get("rename") or {}).items()}
    ign = {norm_name(x) for x in (override.get("ignore") or [])}
    kept = []
    for s in sm.sections:
        k = norm_name(s.name)
        if k in ign:
            continue
        if k in ren:
            s.name = ren[k]
        kept.append(s)
    sm.sections = kept


# ── rendering for the tick-through ─────────────────────────────────────────

def render_text(sm: SheetMap) -> str:
    out = [f"=== {sm.sheet!r}  block={sm.label_col}  source={sm.source}"]
    out.append("  inputs: " + "  ".join(f"{k}={v}({sm.input_values.get(k)})" for k, v in sm.inputs.items()))
    if sm.control_panel:
        out.append("  NOTE: flag block is formula-driven (control panel) — scenario writes override it")
    out.append("  flags:")
    for f in sm.flags:
        out.append(f"    {f.label:24} thickness {f.thickness_cell}={f.thickness!s:8} flag {f.flag_cell}={f.flag}"
                   + ("  (formula)" if f.is_formula else ""))
    out.append(f"  sections ({len(sm.sections)}):")
    for s in sm.sections:
        gates = "; ".join(f"{g.cell}<-{','.join(g.flag_labels)}" for g in s.gate_cells) or "-"
        out.append(f"    {s.name:28} rows {s.first_row}-{s.last_row} hdr={s.header_row} total={s.total_row} "
                   f"cols q/p/t={s.qty_col}/{s.price_col}/{s.total_col} lines={len(s.lines)} "
                   f"gated={','.join(s.gated_cells)} gate={gates} cached={s.cached_total}"
                   + (f"  UNVERIFIABLE:{s.unverifiable}" if s.unverifiable else ""))
    out.append(f"  grand total: {sm.grand_total_cell}={sm.grand_total_cached}  self-check "
               + ("OK" if sm.selfcheck_ok else f"FAILED delta={sm.selfcheck_delta}"))
    if sm.stale_links:
        out.append("  stale external links: " + ", ".join(f"[{k}]={v}" for k, v in sm.stale_links.items()))
    for w in sm.warnings:
        out.append(f"  WARNING: {w}")
    return "\n".join(out)
