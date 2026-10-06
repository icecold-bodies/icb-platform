"""RT5 — Burt's body rules shown in red: the family's rule note (RT5_DISPATCH; RT5_RULING_1).

A family (trailer group) carries a plain-text rule note (migration 0052, trailer_groups.rule_note). Every body of
the family shows it in red under BODY OPTIONS (the calculator) and under its header (Body Templates); the pages
read it from the body's /api/trailers row (`rule_note`, a sibling of RT3's `family`). These tests pin:

  1. the note's red clears 4.6:1 on the MES skin's two backgrounds (the skin's own --red does not), and the stored
     form of a note (line ends, blank edges, the length cap);
  2. the note reaches every body of its family and no other body — the list, the single body (a re-opened costing
     on an inactive body) — and a body with no family shows the OTHER family's;
  3. an admin sets, changes and clears it in the family editor; an older form without the field leaves it alone;
     a note over the cap is refused;
  4. only an admin changes it: 403 for every other role, the note unchanged — the API, not the page, enforces it;
  5. it is stored and served as typed and ESCAPED on the page: an HTML string never becomes markup;
  6. it never reaches a customer document: the report, the Word, the PDF and the Excel are the same with and
     without it, and none carries it;
  7. the two pages carry the note's box in that red, outside the re-rendered lists.

SYNTHETIC rows only (marker RT5F), in the _test database, removed by primary key.
"""
from __future__ import annotations

import io
import json
import uuid
from pathlib import Path

import pytest

MARK = "RT5F"
ROOT = Path(__file__).resolve().parents[2]
NOTE_A = f"{MARK} No PU insulation for Chillers"
NOTE_B = f"{MARK} line one\n{MARK} line two"
XSS = f'<b>{MARK} bold</b> <img src=x onerror="window.__rt5=1"> & "quotes"'
MES_BACKGROUNDS = ("#ffffff", "#f5f7fb")


# ── 1. the red, and the stored form ───────────────────────────────────────────

def test_the_notes_red_reads_on_the_mes_skin_and_the_skins_own_red_does_not():
    from app.services import body_family as bf
    assert bf.RULE_NOTE_INK == "#D12424"
    assert min(bf.contrast(bf.RULE_NOTE_INK, bg) for bg in MES_BACKGROUNDS) >= bf.INK_MIN
    assert bf.worst(bf.SKIN_RED) < bf.INK_MIN          # why the ink is derived, not the skin's --red as-is


@pytest.mark.parametrize("raw,stored", [
    (None, None), ("", None), ("   \r\n  \n", None),
    ("No PU insulation for Chillers", "No PU insulation for Chillers"),
    ("\r\n  first   \r\nsecond\r\n\r\n", "  first\nsecond"),
    ("a\rb", "a\nb"),
])
def test_a_note_is_stored_with_plain_line_ends_and_no_blank_edges(raw, stored):
    from app.services import body_family as bf
    assert bf.normalise_rule_note(raw) == stored


def test_a_note_over_the_cap_is_refused():
    from app.services import body_family as bf
    assert bf.normalise_rule_note("x" * bf.RULE_NOTE_MAX) == "x" * bf.RULE_NOTE_MAX
    with pytest.raises(ValueError):
        bf.normalise_rule_note("x" * (bf.RULE_NOTE_MAX + 1))


# ── fixtures ─────────────────────────────────────────────────────────────────

@pytest.fixture(scope="module")
def app_mod():
    import app.main as m
    from starlette.testclient import TestClient
    with TestClient(m.app):            # startup seeds the admin user + permissions
        yield m


@pytest.fixture(scope="module")
def client(app_mod):
    from starlette.testclient import TestClient
    with TestClient(app_mod.app) as c:
        yield c


def _session(username: str) -> dict:
    from app.database import SessionLocal, User, UserSession
    sid = f"rt5f-{uuid.uuid4().hex[:12]}"
    csrf = f"csrf-{sid}"
    with SessionLocal() as db:
        u = db.query(User).filter_by(username=username).first()
        assert u, f"user {username!r} missing"
        db.merge(UserSession(id=sid, user_id=u.id, role=u.role, expires_at=None, csrf_token=csrf))
        db.commit()
    return {"Cookie": f"session_id={sid}", "X-CSRF-Token": csrf}


@pytest.fixture(scope="module")
def users(app_mod):
    from app.database import SessionLocal, User, UserSession
    made = {}
    with SessionLocal() as db:
        for role in ("sales", "full", "user"):
            uname = f"{MARK.lower()}_{role}_{uuid.uuid4().hex[:6]}"
            db.add(User(username=uname, password_hash="x", role=role))
            made[role] = uname
        db.commit()
    heads = {role: _session(u) for role, u in made.items()}
    heads["admin"] = _session("admin")
    yield heads
    with SessionLocal() as db:
        for h in heads.values():
            db.query(UserSession).filter_by(id=h["Cookie"].split("session_id=")[1]).delete()
        for uname in made.values():
            db.query(User).filter_by(username=uname).delete()
        db.commit()


@pytest.fixture()
def fams(app_mod):
    """ALPHA (a note, and a quote template, so its report prints) with two bodies — one inactive; BETA (no note)
    with one; one body in no family. Every row is removed by primary key, child-first."""
    from app.database import ReportTemplate, SessionLocal, TrailerGroup, TrailerType
    with SessionLocal() as db:
        tmpl = db.query(ReportTemplate).filter_by(is_active=True).order_by(ReportTemplate.id).first()
        ga = TrailerGroup(name=f"{MARK} ALPHA", colour="#168ED9", sort_order=9101, rule_note=NOTE_A,
                          report_template_id=tmpl.id if tmpl else None)
        gb = TrailerGroup(name=f"{MARK} BETA", colour="#4263EB", sort_order=9102)
        db.add_all([ga, gb])
        db.flush()
        ta = TrailerType(name=f"{MARK} ALPHA ONE", is_active=True, group_id=ga.id)
        ta2 = TrailerType(name=f"{MARK} ALPHA TWO (inactive)", is_active=False, group_id=ga.id)
        tb = TrailerType(name=f"{MARK} BETA ONE", is_active=True, group_id=gb.id)
        tn = TrailerType(name=f"{MARK} NO FAMILY", is_active=True)
        db.add_all([ta, ta2, tb, tn])
        db.commit()
        ids = {"ga": ga.id, "gb": gb.id, "ta": ta.id, "ta2": ta2.id, "tb": tb.id, "tn": tn.id}
    yield ids
    with SessionLocal() as db:
        db.query(TrailerType).filter(TrailerType.id.in_([ids["ta"], ids["ta2"], ids["tb"], ids["tn"]])).delete(
            synchronize_session=False)
        db.query(TrailerGroup).filter(TrailerGroup.id.in_([ids["ga"], ids["gb"]])).delete(synchronize_session=False)
        db.commit()


def _note(gid):
    from app.database import SessionLocal, TrailerGroup
    with SessionLocal() as db:
        return db.get(TrailerGroup, gid).rule_note


def _set_note(gid, note):
    from app.database import SessionLocal, TrailerGroup
    with SessionLocal() as db:
        db.get(TrailerGroup, gid).rule_note = note
        db.commit()


def _rows(client, headers, inactive=False) -> dict:
    """The bodies list as that user gets it (inactive bodies are an admin's list)."""
    rows = client.get("/api/trailers" + ("?include_inactive=1" if inactive else ""), headers=headers).json()
    assert isinstance(rows, list), rows
    return {r["id"]: r for r in rows}


def _edit_group(client, headers, gid, **form):
    from app.database import SessionLocal, TrailerGroup
    with SessionLocal() as db:
        g = db.get(TrailerGroup, gid)
        data = {"name": g.name, "description": g.description or "", "report_template_id": g.report_template_id or "",
                "colour": g.colour or "", "sort_order": str(g.sort_order)}
    data.update(form)
    return client.post(f"/admin/quote-templates/groups/{gid}/edit", headers=headers, data=data, follow_redirects=False)


# ── 2. every body of the family, and no other body ────────────────────────────

def test_the_note_reaches_every_body_of_its_family_and_no_other(client, users, fams):
    from app.database import SessionLocal, TrailerGroup
    assert _rows(client, users["sales"])[fams["ta"]]["rule_note"] == NOTE_A   # any signed-in user reads it
    rows = _rows(client, users["admin"], inactive=True)
    assert rows[fams["ta"]]["rule_note"] == NOTE_A
    assert rows[fams["ta2"]]["rule_note"] == NOTE_A             # inactive, still its family's
    assert rows[fams["tb"]]["rule_note"] is None                # BETA has none
    with SessionLocal() as db:                                  # a body in no family shows OTHER's (as its chip does)
        other = db.query(TrailerGroup).filter(TrailerGroup.name.ilike("OTHER")).first()
        other_note = other.rule_note.strip() if other is not None and other.rule_note else None
    assert rows[fams["tn"]]["rule_note"] == (other_note or None)
    carrying = [r["id"] for r in rows.values() if r["rule_note"] == NOTE_A]
    assert sorted(carrying) == sorted([fams["ta"], fams["ta2"]]), "the note leaked to a body outside ALPHA"
    # the family object keeps RT3's ratified shape; the note rides beside it
    assert set(rows[fams["ta"]]["family"]) == {"id", "name", "colour", "ink", "sort_order"}


def test_a_reopened_costing_on_an_inactive_body_gets_the_note_too(client, users, fams):
    r = client.get(f"/api/trailers/{fams['ta2']}", headers=users["sales"])
    assert r.status_code == 200 and r.json()["rule_note"] == NOTE_A


def test_a_new_body_of_the_family_carries_the_note_with_no_code_change(client, users, fams):
    from app.database import SessionLocal, TrailerType
    r = client.post("/api/trailers", headers=users["admin"],
                    json={"name": f"{MARK} ALPHA NEW", "group_id": fams["ga"]})
    assert r.status_code == 200, r.text
    tid = r.json()["id"]
    try:
        assert _rows(client, users["sales"])[tid]["rule_note"] == NOTE_A
    finally:
        with SessionLocal() as db:
            db.query(TrailerType).filter_by(id=tid).delete()
            db.commit()


# ── 3. the admin edits it ─────────────────────────────────────────────────────

def test_an_admin_sets_changes_and_clears_the_note_in_the_family_editor(client, users, fams):
    gid, admin = fams["gb"], users["admin"]
    r = _edit_group(client, admin, gid, rule_note="\r\n" + NOTE_B.replace("\n", "  \r\n") + "\r\n")
    assert r.status_code == 303, r.text
    assert _note(gid) == NOTE_B                                  # stored with plain line ends, no blank edges
    assert _rows(client, users["sales"])[fams["tb"]]["rule_note"] == NOTE_B
    r = _edit_group(client, admin, gid)                          # an older form: no rule_note field at all
    assert r.status_code == 303 and _note(gid) == NOTE_B         # ... leaves the note as it is
    r = _edit_group(client, admin, gid, rule_note="   ")
    assert r.status_code == 303 and _note(gid) is None           # blank clears it
    assert _rows(client, users["sales"])[fams["tb"]]["rule_note"] is None


def test_a_note_over_the_cap_is_refused_and_nothing_changes(client, users, fams):
    r = _edit_group(client, users["admin"], fams["ga"], rule_note="x" * 501)
    assert r.status_code == 400 and "500" in r.text
    assert _note(fams["ga"]) == NOTE_A


def test_a_new_family_can_be_created_with_a_note(client, users, fams):
    from app.database import SessionLocal, TrailerGroup
    r = client.post("/admin/quote-templates/groups/new", headers=users["admin"], follow_redirects=False,
                    data={"name": f"{MARK} GAMMA", "description": "", "colour": "", "sort_order": "9103",
                          "rule_note": NOTE_A})
    assert r.status_code == 303, r.text
    with SessionLocal() as db:
        g = db.query(TrailerGroup).filter_by(name=f"{MARK} GAMMA").first()
        assert g is not None and g.rule_note == NOTE_A
        fams_api = client.get("/api/admin/body-families", headers=users["admin"]).json()
        assert next(f for f in fams_api if f["id"] == g.id)["rule_note"] == NOTE_A
        db.delete(g)
        db.commit()


# ── 4. only an admin ──────────────────────────────────────────────────────────

@pytest.mark.parametrize("role", ["sales", "full", "user"])
def test_a_non_admin_cannot_change_the_note(client, users, fams, role):
    r = _edit_group(client, users[role], fams["ga"], rule_note=f"{MARK} changed by {role}")
    assert r.status_code == 403, r.text
    assert _note(fams["ga"]) == NOTE_A
    r = client.post("/admin/quote-templates/groups/new", headers=users[role], follow_redirects=False,
                    data={"name": f"{MARK} BY {role.upper()}", "rule_note": "x"})
    assert r.status_code == 403
    assert client.get("/api/admin/body-families", headers=users[role]).status_code == 403


# ── 5. escaped, never markup ──────────────────────────────────────────────────

def test_an_html_note_is_stored_as_typed_and_escaped_on_the_family_editor(client, users, fams):
    r = _edit_group(client, users["admin"], fams["ga"], rule_note=XSS)
    assert r.status_code == 303
    assert _note(fams["ga"]) == XSS                                  # stored as typed
    assert _rows(client, users["sales"])[fams["ta"]]["rule_note"] == XSS   # served as JSON text
    html = client.get("/admin/quote-templates", headers=users["admin"]).text
    assert XSS not in html                                           # never as markup ...
    assert f"&lt;b&gt;{MARK} bold&lt;/b&gt;" in html                 # ... but escaped text in the textarea
    assert "onerror=&#34;window.__rt5=1&#34;" in html or "onerror=&quot;window.__rt5=1&quot;" in html


def test_the_pages_write_the_note_as_text_never_as_html():
    """The two pages put the API's string on the screen with textContent only (the journey proves it in a
    browser; this pins the code so an innerHTML can't slip in)."""
    calc = (ROOT / "backend/app/static/js/calculator.js").read_text(encoding="utf-8")
    tmpl = (ROOT / "backend/app/static/js/admin_templates.js").read_text(encoding="utf-8")
    for js, fn in ((calc, "function renderBodyRuleNote()"), (tmpl, "function _showRuleNote(t)")):
        body = js[js.index(fn):js.index("\n}\n", js.index(fn))]
        assert ".textContent = note" in body
        assert "innerHTML" not in body and "insertAdjacentHTML" not in body


# ── 6. never on a customer document ───────────────────────────────────────────

RESULT = {
    "items": [{"category": "FLOOR", "material": "RT5F FLOOR SHEET", "material_code": "R5-1", "formula": "L*W",
               "quantity": 10.0, "unit": "m2", "unit_price": 100.0, "waste_pct": 5, "line_cost": 1050.0,
               "last_updated": None}],
    "category_totals": {"FLOOR": 1050.0}, "cost_per_sqm": 96.0, "grand_total": 1050.0, "profit_margin": 10,
    "profit_amount": 105.0, "ratio_value": 0.55, "ratio_label": "55%", "ratio_amount": 900.0, "selling_price": 2100.0,
}


def _documents(client, headers, rec_id, monkeypatch) -> dict:
    """The four customer documents of one costing, as text. The quote report renders through WeasyPrint (absent
    locally and in CI), so — as test_costing_bom_order does — the renderer is stubbed and EVERYTHING it is handed
    (the template's whole context) is the report's text."""
    from docx import Document
    import openpyxl
    from pypdf import PdfReader
    import app.report_engine as report_engine
    out, seen = {}, {}

    def fake_render(**kw):
        seen.update(kw)
        return b"%PDF-1.4 stub"
    monkeypatch.setattr(report_engine, "render_by_slug", fake_render)
    r = client.get(f"/results/{rec_id}/report", headers=headers)
    assert r.status_code == 200, r.text
    assert seen, "the report renderer was not reached"
    out["report"] = json.dumps(seen, default=str, sort_keys=True)
    r = client.get(f"/results/{rec_id}/export/word", headers=headers)
    assert r.status_code == 200, r.text
    d = Document(io.BytesIO(r.content))
    out["word"] = "|".join([p.text for p in d.paragraphs] + [c.text for t in d.tables for row in t.rows for c in row.cells])
    r = client.get(f"/results/{rec_id}/export/excel", headers=headers)
    assert r.status_code == 200, r.text
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    out["excel"] = "|".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None)
    r = client.get(f"/results/{rec_id}/export/pdf", headers=headers)
    assert r.status_code == 200 and r.content[:5] == b"%PDF-", r.text[:200]
    out["pdf"] = "|".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(r.content)).pages)
    return out


def test_the_note_never_reaches_a_customer_document(client, users, fams, monkeypatch):
    from app.database import CalculationRecord, SessionLocal, User
    with SessionLocal() as db:
        admin = db.query(User).filter_by(username="admin").first()
        rec = CalculationRecord(trailer_type_id=fams["ta"], user_id=admin.id, status="pending", is_repair=False,
                                dimensions_json=json.dumps({"length": 6.0, "width": 2.4, "height": 2.4}),
                                result_json=json.dumps(RESULT))
        db.add(rec)
        db.commit()
        rec_id = rec.id
    try:
        with_note = _documents(client, users["admin"], rec_id, monkeypatch)
        for name, text in with_note.items():
            assert MARK + " No PU" not in text, f"the rule note reached the {name}"
        _set_note(fams["ga"], None)
        without = _documents(client, users["admin"], rec_id, monkeypatch)
        for name in ("word", "excel", "pdf"):
            assert with_note[name] == without[name], f"the {name} changed with the note"
        assert with_note["report"] == without["report"]
    finally:
        with SessionLocal() as db:
            db.query(CalculationRecord).filter_by(id=rec_id).delete()
            db.commit()


# ── 7. the two pages carry the box, in the note's red ─────────────────────────

def test_the_calculator_and_body_templates_carry_the_note_box_in_its_red(client, users):
    from app.services import body_family as bf
    for path, box in (("/mes/calculator?stay=1", 'id="body-rule-note"'), ("/admin/templates", 'id="tt-rule-note"')):
        html = client.get(path, headers=users["admin"]).text
        assert box in html, path
        i = html.index(box)
        tag = html[html.rindex("<div", 0, i):html.index(">", i)]
        assert f"color:{bf.RULE_NOTE_INK}" in tag, (path, tag)
        assert "hidden" in tag, "the box starts hidden; the page shows it only for a body with a note"
    calc = client.get("/mes/calculator?stay=1", headers=users["admin"]).text
    # under the BODY OPTIONS heading, outside the re-rendered list
    assert calc.index("Body Options") < calc.index('id="body-rule-note"') < calc.index('id="body-options-list"')


def test_no_customer_document_code_reads_the_family_rule_note():
    """Belt and braces for (6): no exporter, report engine, quote-document builder or document template reads the
    family's rule note. (`zero_rule_note` is v1.44's own document note, a different thing.)"""
    import re
    rx = re.compile(r"(?<![A-Za-z_])rule_note(?!_ink)")
    app = ROOT / "backend" / "app"
    files = [app / "routers" / "exports.py", app / "report_engine.py", app / "pdf_generator.py",
             app / "services" / "quote_document.py", app / "services" / "document_context.py"]
    files += sorted((app / "templates" / "reports").rglob("*")) + sorted((app / "templates" / "pdf_templates").rglob("*"))
    files += sorted((app / "pdf_templates").rglob("*"))
    hits = [f"{f.relative_to(ROOT)}:{n}" for f in files if f.is_file() and f.suffix in (".py", ".html", ".j2", ".txt")
            for n, line in enumerate(f.read_text(encoding="utf-8", errors="replace").splitlines(), 1) if rx.search(line)]
    assert not hits, f"customer-document code reads the rule note: {hits}"
