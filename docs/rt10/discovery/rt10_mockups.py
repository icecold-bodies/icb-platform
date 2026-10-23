"""RT10 §3.0 — MOCK-UPS (nothing here is the build): the proposed REV display, drawn on today's real surfaces.

  documents (in-process, today's renderers): the costing export (PDF page 1; Word + Excel heading text) and the
  letterhead quote ("Generate Quote" HTML template, screenshot in Chromium because WeasyPrint is not installed
  here) — each at REV 0 (unchanged) and REV 1 (the proposed label);
  staff screens (the :8012 side server, today's SPA/calculator, the proposed labels injected into the DOM): the
  Costings list group, the costing view header + Revisions list (latest and older), the calculator's edit banner
  on an older revision.

Marker rows RT10J* on icb_test only (purged at the start and the end). Usage (from the rt10 backend dir):
  DATABASE_URL=...icb_test python rt10_mockups.py <out_dir>"""
from __future__ import annotations

import copy
import json
import os
import sys
from pathlib import Path

from playwright.sync_api import expect, sync_playwright

BASE = os.environ.get("MES_BASE", "http://127.0.0.1:8012").rstrip("/")
assert ":8000" not in BASE and os.environ.get("DATABASE_URL", "").endswith("/icb_test")
OUT = Path(sys.argv[1])
OUT.mkdir(parents=True, exist_ok=True)
sys.path.insert(0, os.getcwd())
from app.database import CalculationRecord, Customer, SessionLocal, TrailerType, UserSession  # noqa: E402
from sqlalchemy import text  # noqa: E402

T = 20_000
SNAP = {"body_option_selections": {}, "drd_srd": {}, "draft_flag_state": {}, "draft_flag_vars": {},
        "draft_category_radio": {}, "draft_masterless_cat": {}, "draft_folder": {}, "optional_sections_enabled": [],
        "optional_row_excl": [], "price_overrides": {}, "ratio": "0.55", "chassis": {"enabled": False},
        "body_variables": {}}


def purge() -> None:
    with SessionLocal() as db:
        db.execute(text("DELETE FROM icb_costings.calculations c USING icb_costings.customers cu "
                        "WHERE c.customer_id = cu.id AND cu.name LIKE 'RT10J%'"))
        db.execute(text("DELETE FROM icb_costings.customers WHERE name LIKE 'RT10J%'"))
        db.commit()


def csrf(page) -> str:
    sid = next((c["value"] for c in page.context.cookies() if c["name"] == "session_id"), None)
    with SessionLocal() as db:
        row = db.get(UserSession, sid)
        return (row.csrf_token or "") if row else ""


def approve(page, body: dict) -> dict:
    r = page.request.post(f"{BASE}/api/approve", data=json.dumps(body),
                          headers={"Origin": BASE, "X-CSRF-Token": csrf(page), "Content-Type": "application/json"})
    assert r.ok, r.text()
    return r.json()


def _dmy(s: str) -> str:
    from datetime import datetime
    return datetime.strptime(s, "%Y-%m-%d %H:%M").strftime("%d %b %Y %H:%M").lstrip("0")


def rev_label(version) -> str:
    return f"REV {max(int(version or 1) - 1, 0)}"


# ── documents ────────────────────────────────────────────────────────────────
def documents(rec_id: int, rev1_id: int) -> dict:
    import fitz
    from app.routers import exports as ex
    from app import report_engine as re_
    notes = {}
    with SessionLocal() as db:
        rec = db.query(CalculationRecord).filter_by(id=rec_id).first()
        ctx0, _stem = ex._doc_ctx_for_record(rec, db, detail="items", ratios_raw=None)
        qn = rec.quote_number
        ctx1 = copy.deepcopy(ctx0)
        ctx1["heading"] = ctx0["heading"].replace(qn, f"{qn} · REV 1", 1)     # the proposed label, REV 1 only
        for label, ctx in (("rev0", ctx0), ("rev1", ctx1)):
            pdf = ex._render_pdf(ctx)
            doc = fitz.open(stream=pdf, filetype="pdf")
            pix = doc[0].get_pixmap(matrix=fitz.Matrix(1.6, 1.6), clip=fitz.Rect(0, 0, doc[0].rect.width, 230))
            pix.save(str(OUT / f"doc_export_pdf_{label}.png"))
            subject, _body = ex._email_body(ctx, mode="approved", sender="admin", note="")
            notes[label] = {"heading": ctx["heading"], "email_subject": subject}
            # Word + Excel: the same ctx heading is their first line
            from docx import Document
            d = Document(ex._render_docx(ctx))
            notes[label]["word_first_line"] = next((p.text for p in d.paragraphs if p.text.strip()), "")
            from openpyxl import load_workbook
            wb = load_workbook(ex._render_xlsx(ctx))
            ws = wb.worksheets[0]
            notes[label]["excel_first_cell"] = next((str(c.value) for row in ws.iter_rows(max_row=6) for c in row
                                                     if c.value not in (None, "")), "")
        # the letterhead quote ("Generate Quote" → templates/reports/<slug>.html), REV 0 and REV 1
        tt = db.query(TrailerType).filter_by(id=rec.trailer_type_id).first()
        from app.services import resolve_report_template
        slug = resolve_report_template(tt).slug
        result = json.loads(rec.result_json)
        cust = {"name": "RT10J MOCKUP CUSTOMER", "email": "", "telephone": "", "contact_name": "", "contact_email": "",
                "contact_telephone": ""}
        builder = re_.CONTEXT_BUILDERS.get(slug, re_.build_explosive_quote_context)
        for label, ref in (("rev0", qn), ("rev1", f"{qn} · REV 1")):
            lctx = builder(record_id=rec.id, customer=cust, dimensions=json.loads(rec.dimensions_json),
                           result=result, created_at=rec.created_at, quote_number=ref)
            html = re_._env.get_template(f"{slug}.html").render(**lctx)
            p = OUT / f"_letterhead_{label}.html"
            p.write_text(html, encoding="utf-8")
            notes[f"letterhead_{label}"] = {"slug": slug, "ref_no": ref}
    return notes


def shoot_letterhead(browser) -> None:
    page = browser.new_page(viewport={"width": 900, "height": 1300})
    for label in ("rev0", "rev1"):
        page.goto((OUT / f"_letterhead_{label}.html").resolve().as_uri())
        page.wait_for_timeout(800)
        page.screenshot(path=str(OUT / f"doc_letterhead_{label}_top.png"), clip={"x": 0, "y": 0, "width": 900, "height": 520})
        acc = page.locator("text=ACCEPTANCE OF QUOTE NO").first
        if acc.count():
            acc.scroll_into_view_if_needed()
            box = acc.bounding_box()
            page.screenshot(path=str(OUT / f"doc_letterhead_{label}_acceptance.png"),
                            clip={"x": 0, "y": max(box["y"] - 40, 0), "width": 900, "height": 160})
    page.close()


# ── staff screens ────────────────────────────────────────────────────────────
INJECT_BOARD = """(info) => {
  const css = (el, s) => { el.setAttribute('style', s); return el; };
  const rows = [...document.querySelectorAll("[data-testid='costing-row']")];
  const byId = new Map(rows.map(tr => [+tr.querySelector("a[href*='/costings/results/']").getAttribute('href').split('/').pop(), tr]));
  const tbody = rows[0].parentNode;
  // the group sits where its latest falls: move the older revisions straight under it, latest on top
  let anchor = byId.get(info.order[0]);
  for (const id of info.order.slice(1)) { const tr = byId.get(id); anchor.after(tr); anchor = tr; }
  for (const r of info.revs) {
    const tr = byId.get(r.id), cell = tr.querySelector('td:nth-child(2)');
    cell.style.whiteSpace = 'nowrap';
    const chip = document.createElement('span');
    if (r.latest) {
      chip.innerHTML = '<span style="font-weight:700">· REV ' + r.rev + '</span>'
        + ' <span style="margin-left:4px;border-radius:9999px;background:#16a34a22;color:#15803d;padding:1px 7px;font-size:10px;font-weight:700;font-family:Inter,sans-serif">Latest</span>'
        + '<div style="font-family:Inter,sans-serif;font-weight:400;font-size:10px;color:#6b7280;margin-top:2px">' + info.revs.length + ' revisions</div>';
    } else {
      chip.innerHTML = '<span style="color:#9ca3af;font-weight:600">· REV ' + r.rev + ' · earlier</span>';
      tr.style.opacity = '0.8';
    }
    cell.appendChild(chip);
  }
  for (const id of info.solo) {   // a quote with no revisions: REV 0, no Latest marker
    const cell = byId.get(id).querySelector('td:nth-child(2)');
    cell.style.whiteSpace = 'nowrap';
    const s = document.createElement('span'); s.innerHTML = '<span style="font-weight:700">· REV 0</span>'; cell.appendChild(s);
  }
}"""

INJECT_HEADER = """(info) => {
  const h1 = document.querySelector('h1'), q = h1.querySelector('span.font-mono');
  const add = document.createElement('span');
  add.setAttribute('style', 'font-family:ui-monospace,monospace');
  add.innerHTML = info.latest
    ? ' · REV ' + info.rev + ' <span style="margin-left:6px;border-radius:9999px;background:#16a34a22;color:#15803d;padding:2px 9px;font-size:12px;font-weight:700;font-family:Inter,sans-serif">Latest</span>'
    : ' · REV ' + info.rev + ' <a href="#" style="margin-left:8px;font-family:Inter,sans-serif;font-size:13px;font-weight:600;color:#2563eb;text-decoration:underline">REV ' + info.latestRev + ' is the latest →</a>';
  q.after(add);
  const card = document.createElement('div');
  card.setAttribute('style', 'margin:10px 0 14px;max-width:760px;border:1px solid #e5e7eb;border-radius:8px;background:#fff;padding:10px 14px;font-family:Inter,sans-serif;font-size:13px');
  let html = '<div style="font-weight:700;margin-bottom:6px">Revisions</div>';
  if (info.repriced) html += '<div style="margin-bottom:8px;color:#b45309">' + info.repriced + '</div>';
  html += '<table style="width:100%;border-collapse:collapse"><tr style="color:#6b7280;text-align:left;font-size:11px"><th style="padding:2px 12px 2px 0">REV</th><th style="padding:2px 12px">Saved</th><th style="text-align:right;padding:2px 12px">Net total</th><th style="padding:2px 12px">Status</th><th></th></tr>';
  for (const r of info.list) {
    const me = r.id === info.me;
    html += '<tr style="border-top:1px solid #f3f4f6;' + (me ? 'background:#eff6ff' : '') + '"><td style="padding:4px 0;font-family:ui-monospace,monospace;font-weight:700">REV ' + r.rev + (r.latest ? ' <span style="border-radius:9999px;background:#16a34a22;color:#15803d;padding:0 6px;font-size:10px;font-family:Inter,sans-serif">Latest</span>' : '') + '</td><td style="padding:4px 12px">' + r.saved + '</td><td style="text-align:right;padding:4px 12px;font-variant-numeric:tabular-nums">' + r.net + '</td><td style="padding:4px 12px">' + r.status + '</td><td style="text-align:right">' + (me ? '<span style="color:#6b7280">this one</span>' : '<a href="#" style="color:#2563eb">Open</a>') + '</td></tr>';
  }
  html += '</table>';
  card.innerHTML = html;
  h1.parentNode.parentNode.after(card);
}"""

INJECT_BANNER = """(info) => {
  const t = document.getElementById('edit-banner-text');
  t.innerHTML = t.innerHTML.replace(/· Rev\\d+/, '· REV ' + info.rev);
  const line = document.createElement('div');
  line.setAttribute('style', 'margin-top:4px;font-size:12px;color:#b45309');
  line.innerHTML = 'REV ' + info.latestRev + ' is the latest. <a href="#" style="color:#2563eb;font-weight:600;text-decoration:underline">Open it instead?</a>';
  t.appendChild(line);
}"""


def main() -> None:
    purge()
    with SessionLocal() as db:
        c1, c2 = Customer(name="RT10J MOCKUP CUSTOMER"), Customer(name="RT10J OTHER CUSTOMER")
        db.add_all([c1, c2])
        db.commit()
        cust, cust2 = c1.id, c2.id
        tt = db.query(TrailerType).filter_by(name="FREEZER MEDIUM").first()
        tt_id, dims = tt.id, {"length": tt.default_length or 6.0, "width": tt.default_width or 2.5,
                              "height": tt.default_height or 2.4}
    base = {"trailer_type_id": tt_id, "dimensions": dims, "profit_margin": 5, "overrides": {}, "override_reasons": {},
            "chassis": {"enabled": False}, "user_excluded_bom_ids": [], "optional_sections_enabled": [],
            "ratio_value": 0.55, "ratio_label": "55%", "ui_snapshot": SNAP}
    with sync_playwright() as p:
        browser = p.chromium.launch()
        ctx = browser.new_context(base_url=BASE, viewport={"width": 1440, "height": 900})
        page = ctx.new_page()
        page.set_default_timeout(T)
        assert page.request.post(f"{BASE}/api/mes/autologin", headers={"Origin": BASE}).ok
        v1 = approve(page, {**base, "customer_id": cust, "version_action": "save_as_new"})
        v2 = approve(page, {**base, "customer_id": cust, "version_action": "new_version", "next_version": 2,
                            "reuse_quote_number": True, "edit_record_id": v1["record_id"]})
        other = approve(page, {**base, "customer_id": cust2, "version_action": "save_as_new"})
        v3 = approve(page, {**base, "customer_id": cust, "version_action": "new_version", "next_version": 3,
                            "reuse_quote_number": True, "edit_record_id": v1["record_id"],
                            "profit_margin": 6})
        qn = v1["quote_number"]
        ids = {"v1": v1["record_id"], "v2": v2["record_id"], "other": other["record_id"], "v3": v3["record_id"]}
        notes = {"quote_number": qn, "ids": ids, "other_quote_number": other["quote_number"]}
        notes["documents"] = documents(ids["v1"], ids["v2"])
        shoot_letterhead(browser)

        # the board: today, then the proposal
        page.goto(f"{BASE}/mes-app/costings")
        page.wait_for_selector("[data-testid='costings-table']")
        expect(page.locator("[data-testid='costing-row']").first).to_be_visible()
        page.wait_for_timeout(1500)
        table = page.locator("[data-testid='costings-table']")
        box = table.bounding_box()
        clip = {"x": box["x"], "y": box["y"], "width": min(box["width"], 1400), "height": 330}
        page.screenshot(path=str(OUT / "screen_list_today.png"), clip=clip)
        page.evaluate(INJECT_BOARD, {"order": [ids["v3"], ids["v2"], ids["v1"]],
                                     "revs": [{"id": ids["v3"], "rev": 2, "latest": True},
                                              {"id": ids["v2"], "rev": 1, "latest": False},
                                              {"id": ids["v1"], "rev": 0, "latest": False}],
                                     "solo": [ids["other"]]})
        page.screenshot(path=str(OUT / "screen_list_proposed.png"), clip=clip)

        # the costing view: header + Revisions list (latest, then an older one)
        listing = {c["id"]: c for c in page.request.get(f"{BASE}/api/calculations?limit=100").json()}
        rows_ = [{"id": ids[k], "rev": r, "latest": k == "v3", "saved": _dmy(listing[ids[k]]["created_at"]),
                  "net": f"R {listing[ids[k]]['grand_total']:,.2f}".replace(",", " "), "status": listing[ids[k]]["mes_status"]}
                 for k, r in (("v3", 2), ("v2", 1), ("v1", 0))]
        for label, me, rev, latest, rp in (
                ("latest", ids["v3"], 2, True, "Moved to today's prices on 22 Oct 2026, from REV 0"),
                ("older", ids["v1"], 0, False, None)):
            page.goto(f"{BASE}/mes-app/costings/{qn.replace('/', '%2F')}")
            page.wait_for_selector("h1 span.font-mono")
            page.wait_for_timeout(1200)
            page.evaluate(INJECT_HEADER, {"latest": latest, "rev": rev, "latestRev": 2, "me": me, "list": rows_,
                                          "repriced": rp})
            page.screenshot(path=str(OUT / f"screen_view_{label}.png"), clip={"x": 0, "y": 40, "width": 1100, "height": 330})

        # the calculator's edit banner on an older revision
        page.goto(f"{BASE}/mes/calculator?edit={ids['v1']}&stay=1")
        expect(page.locator("#edit-mode-banner")).to_be_visible()
        page.wait_for_timeout(1500)
        page.locator("#edit-mode-banner").screenshot(path=str(OUT / "screen_banner_today.png"))
        page.evaluate(INJECT_BANNER, {"rev": 0, "latestRev": 2})
        page.locator("#edit-mode-banner").screenshot(path=str(OUT / "screen_banner_proposed.png"))
        browser.close()
    (OUT / "mockup_notes.json").write_text(json.dumps(notes, indent=1, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(notes, indent=1, ensure_ascii=False))
    purge()
    print("purged RT10J* rows")


if __name__ == "__main__":
    main()
