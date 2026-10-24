"""RT10 — revisions visible: REV n everywhere, the Latest marker, the server numbering, Replace as a soft delete.

Michael's rule (22 Oct 2026): REV = the stored version − 1 (the first save is REV 0); the stored version is never
rewritten. RT10_RULING_1: the quote number is the chain; Save as Rev. keeps it (server-enforced) and takes
max(version over the whole chain, deleted rows included) + 1, decided by the server under the chain's row lock — the
client's next_version is ignored; Latest = the highest version among live rows (tie → newer, then higher id); the list
and the costing GET carry version / rev / revision_count / latest_version / latest_rev / latest_id / is_latest,
computed once per request (no N+1); Replace soft-deletes and names what it would replace; "Generate Quote" is withheld
("Official quotes are issued from SAP.", RULING_1A); the documents carry `· REV n` from REV 1 (RULING_1 Q4).

Negative control (F6): this file is written to run UNCHANGED against the base code (dce981e) — the API-level tests
must fail there for the right reason (2, 2 numbering; a new quote number; no chain fields; a hard-deleted Replace;
the Generate Quote link; no REV on the documents). The helper tests import the new module lazily.

House pattern: live test DB (the db-guard aborts on anything but *_test), marker rows (RT10R), real session rows,
purge child-first by id at setup and teardown.
"""
from __future__ import annotations

import io
import json
import uuid

import pytest
from sqlalchemy import event, text

MARK = "RT10R"
DIMS = {"length": 6.0, "width": 2.4, "height": 2.4, "floor_thickness": 0.1, "panel_thickness": 0.06,
        "insulation_thickness": 0.06, "num_axles": 2, "num_doors": 2}


# ── fixtures ─────────────────────────────────────────────────────────────────
@pytest.fixture(scope="module")
def app_mod():
    import app.main as m
    from starlette.testclient import TestClient
    with TestClient(m.app):
        yield m


@pytest.fixture(scope="module")
def client(app_mod):
    from starlette.testclient import TestClient
    with TestClient(app_mod.app) as c:
        yield c


@pytest.fixture(scope="module")
def admin(app_mod):
    from app.database import SessionLocal, User, UserSession
    sid = f"rt10r-{uuid.uuid4().hex[:12]}"
    with SessionLocal() as db:
        u = db.query(User).filter_by(username="admin").first()
        db.merge(UserSession(id=sid, user_id=u.id, role=u.role, expires_at=None, csrf_token=f"csrf-{sid}"))
        db.commit()
    yield {"Cookie": f"session_id={sid}", "X-CSRF-Token": f"csrf-{sid}"}
    with SessionLocal() as db:
        db.query(UserSession).filter_by(id=sid).delete()
        db.commit()


def _purge() -> None:
    """Marker rows, child-first by id (never TRUNCATE)."""
    from app.database import BillOfMaterial, CalculationRecord, Customer, Material, SessionLocal, TrailerType
    with SessionLocal() as db:
        tids = [t.id for t in db.query(TrailerType).filter(TrailerType.name.like(f"{MARK}%")).all()]
        cids = [c.id for c in db.query(Customer).filter(Customer.name.like(f"{MARK}%")).all()]
        calc_ids = [r.id for r in db.query(CalculationRecord).filter(
            (CalculationRecord.trailer_type_id.in_(tids or [-1])) | (CalculationRecord.customer_id.in_(cids or [-1])))]
        if calc_ids:
            db.execute(text("DELETE FROM icb_mes.production_jobs WHERE calculation_record_id = ANY(:i)"), {"i": calc_ids})
            db.query(CalculationRecord).filter(CalculationRecord.id.in_(calc_ids)).delete(synchronize_session=False)
        if tids:
            db.query(BillOfMaterial).filter(BillOfMaterial.trailer_type_id.in_(tids)).delete(synchronize_session=False)
            db.query(TrailerType).filter(TrailerType.id.in_(tids)).delete(synchronize_session=False)
        db.query(Material).filter_by(material_code=MARK).delete(synchronize_session=False)
        if cids:
            db.query(Customer).filter(Customer.id.in_(cids)).delete(synchronize_session=False)
        db.commit()


@pytest.fixture(scope="module")
def ids(app_mod):
    from app.database import BillOfMaterial, Customer, Material, ReportTemplate, SessionLocal, TrailerType
    _purge()
    out = {}
    with SessionLocal() as db:
        c1, c2, c3 = (Customer(name=f"{MARK} CUSTOMER"), Customer(name=f"{MARK} OTHER"),
                      Customer(name=f"{MARK} REPLACE"))
        tmpl = db.query(ReportTemplate).filter_by(is_active=True).order_by(ReportTemplate.id).first()
        tt = TrailerType(name=f"{MARK} BODY", is_active=True, default_length=6.0, default_width=2.4, default_height=2.4,
                         override_report_template_id=tmpl.id if tmpl else None)
        db.add_all([c1, c2, c3, tt])
        db.flush()
        for i, (name, price, formula) in enumerate((("PANEL", 250.0, "length*width"), ("FLOOR", 300.0, "length*width"))):
            m = Material(name=f"{MARK} {name}", unit_of_measure="m²", price_per_unit=price, material_code=MARK,
                         is_active=True)
            db.add(m)
            db.flush()
            db.add(BillOfMaterial(trailer_type_id=tt.id, material_id=m.id, formula_expression=formula,
                                  waste_percentage=0.0, bom_section=name, sort_order=10 * (i + 1)))
        db.commit()
        out.update(customer=c1.id, other=c2.id, replace=c3.id, tt=tt.id, has_template=tmpl is not None)
    yield out
    _purge()


def _payload(ids, customer_key="customer", **over) -> dict:
    p = {"trailer_type_id": ids["tt"], "dimensions": dict(DIMS), "profit_margin": 5, "overrides": {},
         "override_reasons": {}, "chassis": {"enabled": False}, "user_excluded_bom_ids": [],
         "optional_sections_enabled": [], "customer_id": ids[customer_key], "ratio_value": 0.55,
         "ratio_label": "55%", "ui_snapshot": {"ratio": "0.55"}}
    p.update(over)
    return p


def _save(client, admin, body) -> dict:
    r = client.post("/api/approve", json=body, headers=admin)
    assert r.status_code == 200, r.text
    return r.json()


def _version(record_id: int) -> int:
    from app.database import CalculationRecord, SessionLocal
    with SessionLocal() as db:
        return int(json.loads(db.query(CalculationRecord).filter_by(id=record_id).first().result_json).get("version", 1))


def _qn(record_id: int):
    from app.database import CalculationRecord, SessionLocal
    with SessionLocal() as db:
        return db.query(CalculationRecord).filter_by(id=record_id).first().quote_number


@pytest.fixture(scope="module")
def chain(client, admin, ids):
    """REV 0 and two Save as Rev. from REV 0, each sending what the old client sent (next_version 2, reuse ticked)."""
    first = _save(client, admin, _payload(ids, version_action="save_as_new"))
    revs = [_save(client, admin, _payload(ids, version_action="new_version", next_version=2,
                                          reuse_quote_number=True, edit_record_id=first["record_id"]))
            for _ in range(2)]
    return {"first": first["record_id"], "qn": first["quote_number"], "r1": revs[0]["record_id"],
            "r2": revs[1]["record_id"]}


# ── 1. the mapping (the new helper; imported lazily: absent on the base code) ─
def test_rev_is_version_minus_one_and_the_labels_follow_the_two_formats():
    from app.services import revisions as rv
    assert [rv.rev_number(v) for v in (1, 2, 3, None, 0, "2")] == [0, 1, 2, 0, 0, 1]
    assert rv.version_of('{"version": 3}') == 3 and rv.version_of('{"items": []}') == 1 and rv.version_of("{bad") == 1
    assert rv.staff_label("N10022/10/2026", 1) == "N10022/10/2026 · REV 0"        # staff screens: REV 0 shown
    assert rv.staff_label("N10022/10/2026", 3) == "N10022/10/2026 · REV 2"
    assert rv.doc_label("N10022/10/2026", 1) == "N10022/10/2026"                  # documents: none on REV 0
    assert rv.doc_label("N10022/10/2026", 2) == "N10022/10/2026 · REV 1"
    assert rv.staff_label(None, 1, 42) == "#42 · REV 0"


# ── 2. numbering ─────────────────────────────────────────────────────────────
def test_two_saves_from_rev_0_give_rev_1_and_rev_2_under_the_same_number(chain):
    # 22 Oct: three saves from REV 0 wrote version 2 three times (the client sent editingVersion + 1)
    assert [_version(chain["first"]), _version(chain["r1"]), _version(chain["r2"])] == [1, 2, 3]
    assert _qn(chain["r1"]) == _qn(chain["r2"]) == chain["qn"]


def test_a_revision_keeps_its_number_even_when_the_client_unticks_reuse(client, admin, ids, chain):
    r = _save(client, admin, _payload(ids, version_action="new_version", next_version=2, reuse_quote_number=False,
                                      edit_record_id=chain["first"]))
    assert r["quote_number"] == chain["qn"] and _version(r["record_id"]) == 4
    chain["r3"] = r["record_id"]


def test_a_number_is_never_reused_after_a_delete(client, admin, ids, chain):
    r = client.delete(f"/api/calculations/{chain['r3']}", headers=admin)
    assert r.status_code == 200, r.text
    nxt = _save(client, admin, _payload(ids, version_action="new_version", next_version=2,
                                        edit_record_id=chain["first"]))
    assert _version(nxt["record_id"]) == 5, "REV 3 was issued (then deleted): the next is REV 4, never REV 3 again"
    chain["r4"] = nxt["record_id"]


def test_the_duplicate_path_revises_the_quote_it_names_and_refuses_without_one(client, admin, ids, chain):
    r = _save(client, admin, _payload(ids, version_action="new_version", next_version=2,
                                      revision_of_record_id=chain["r4"]))
    assert r["quote_number"] == chain["qn"] and _version(r["record_id"]) == 6
    chain["r5"] = r["record_id"]
    r = client.post("/api/approve", json=_payload(ids, version_action="new_version", next_version=2),
                    headers=admin)
    assert r.status_code == 409 and "needs the quote it revises" in r.json()["detail"]
    r = client.post("/api/approve", json=_payload(ids, "other", version_action="new_version", next_version=2,
                                                  revision_of_record_id=chain["first"]), headers=admin)
    assert r.status_code == 409 and "another customer or body type" in r.json()["detail"]


def test_a_repair_saved_as_a_new_revision_is_still_a_new_repair_quote(client, admin, ids):
    # RT10_RULING_1 Q2 (a): unchanged — the modal now says what the server does
    rep = {"is_repair": True, "trailer_type_id": None, "repair_type": f"{MARK} door", "repair_scope": None,
           "repair_lines": [{"kind": "free_hand", "key": "f1", "description": "Labour", "qty": 1, "unit": "each",
                             "unit_price": 500, "excluded": False}],
           "profit_margin": 5, "ratio_value": 0.55, "ratio_label": "55%", "repair_vehicle": None, "dimensions": {},
           "customer_id": ids["customer"]}
    a = _save(client, admin, rep)
    b = _save(client, admin, {**rep, "version_action": "new_version", "edit_record_id": a["record_id"]})
    assert b["quote_number"] != a["quote_number"] and _version(b["record_id"]) == 1


# ── 3. the fields, once per request ──────────────────────────────────────────
def _rows(client, admin, **params):
    r = client.get("/api/calculations", params={"limit": 200, **params}, headers=admin)
    assert r.status_code == 200
    return {c["id"]: c for c in r.json()}


def test_the_list_carries_rev_and_the_chain_and_only_the_latest_live_row_is_latest(client, admin, chain):
    rows = _rows(client, admin)
    mine = {k: rows[chain[k]] for k in ("first", "r1", "r2", "r4", "r5")}
    assert chain["r3"] not in rows                                           # deleted: off the board
    assert {k: (r["version"], r["rev"]) for k, r in mine.items()} == {
        "first": (1, 0), "r1": (2, 1), "r2": (3, 2), "r4": (5, 4), "r5": (6, 5)}
    assert {r["revision_count"] for r in mine.values()} == {5}              # live rows only
    assert {(r["latest_version"], r["latest_rev"], r["latest_id"]) for r in mine.values()} == {(6, 5, chain["r5"])}
    assert [k for k, r in mine.items() if r["is_latest"]] == ["r5"]
    deleted = _rows(client, admin, filter="deleted")
    assert deleted[chain["r3"]]["rev"] == 3 and deleted[chain["r3"]]["is_latest"] is False
    assert deleted[chain["r3"]]["latest_id"] == chain["r5"]


def test_the_chain_fields_cost_one_query_per_request_not_one_per_row(client, admin, ids, chain):
    from app.database import engine
    seen = []

    def count(conn, cursor, statement, params, context, executemany):
        if "calculations.quote_number IN" in statement:
            seen.append(statement)
    # a second chain on the page, so a per-row implementation would show itself
    a = _save(client, admin, _payload(ids, "other", version_action="save_as_new"))
    _save(client, admin, _payload(ids, "other", version_action="new_version", edit_record_id=a["record_id"]))
    event.listen(engine, "before_cursor_execute", count)
    try:
        rows = _rows(client, admin)
    finally:
        event.remove(engine, "before_cursor_execute", count)
    assert sum(1 for r in rows.values() if r["revision_count"] > 1) >= 7
    assert len(seen) == 1, f"{len(seen)} chain queries for one list request"


def test_negative_control_a_per_row_chain_lookup_is_counted(client, admin, chain, monkeypatch):
    from app.database import engine
    from app.services import revisions as rv
    real = rv.chain_info
    monkeypatch.setattr(rv, "chain_info", lambda db, records: {k: v for r in records for k, v in real(db, [r]).items()})
    seen = []

    def count(conn, cursor, statement, params, context, executemany):
        if "calculations.quote_number IN" in statement:
            seen.append(statement)
    event.listen(engine, "before_cursor_execute", count)
    try:
        _rows(client, admin)
    finally:
        event.remove(engine, "before_cursor_execute", count)
    assert len(seen) > 1                                       # the counter sees an N+1 when there is one


def test_the_costing_get_carries_the_chain_and_the_revisions_list_with_the_deleted_one_marked(client, admin, chain):
    r = client.get(f"/api/calculations/{chain['first']}", headers=admin)
    assert r.status_code == 200
    d = r.json()
    assert (d["version"], d["rev"], d["is_latest"], d["latest_rev"], d["latest_id"]) == (1, 0, False, 5, chain["r5"])
    revs = d["revisions"]
    assert [x["rev"] for x in revs] == [5, 4, 3, 2, 1, 0]                     # latest first
    assert [x["id"] for x in revs if x["deleted"]] == [chain["r3"]]
    assert [x["id"] for x in revs if x["is_latest"]] == [chain["r5"]]
    assert all(set(x) >= {"id", "rev", "saved_at", "net_total", "status"} for x in revs)


def test_the_duplicate_check_groups_by_quote_and_names_the_next_rev(client, admin, ids, chain):
    r = client.get("/api/check-duplicate", params={"customer_id": ids["customer"], "trailer_type_id": ids["tt"]},
                   headers=admin)
    d = r.json()
    q = [x for x in d["quotes"] if x["quote_number"] == chain["qn"]]
    assert len(q) == 1 and (q[0]["latest_id"], q[0]["latest_rev"], q[0]["next_rev"]) == (chain["r5"], 5, 6)


# ── 4. Replace is a soft delete ──────────────────────────────────────────────
def test_replace_soft_deletes_the_replaced_costings_and_keeps_them_restorable(client, admin, ids):
    from app.database import CalculationRecord, SessionLocal
    a = _save(client, admin, _payload(ids, "replace", version_action="save_as_new"))
    b = _save(client, admin, _payload(ids, "replace", version_action="new_version", edit_record_id=a["record_id"]))
    new = _save(client, admin, _payload(ids, "replace", version_action="replace"))
    with SessionLocal() as db:
        gone = {r.id: r for r in db.query(CalculationRecord).filter(
            CalculationRecord.id.in_([a["record_id"], b["record_id"]])).all()}
        assert set(gone) == {a["record_id"], b["record_id"]}, "Replace hard-deleted the costings"
        assert all(r.deleted_at is not None and r.deleted_by == "admin" for r in gone.values())
    assert _version(new["record_id"]) == 1 and new["quote_number"] != a["quote_number"]
    assert client.post(f"/api/calculations/{a['record_id']}/restore", json={}, headers=admin).status_code == 200


def test_replace_refuses_and_changes_nothing_when_a_costing_has_a_production_job(client, admin, ids):
    from app.database import CalculationRecord, SessionLocal
    rows = [r for r in _rows(client, admin).values() if r["customer"] == f"{MARK} REPLACE"]
    assert rows, "the previous test restored one"
    target = rows[0]["id"]
    with SessionLocal() as db:
        branch = db.execute(text("SELECT id FROM icb_costings.branches ORDER BY id LIMIT 1")).scalar()
        job = db.execute(text("INSERT INTO icb_mes.production_jobs (calculation_record_id, branch_id, status) "
                              "VALUES (:c, :b, 'accepted') RETURNING id"), {"c": target, "b": branch}).scalar()
        db.commit()
    try:
        r = client.post("/api/approve", json=_payload(ids, "replace", version_action="replace"), headers=admin)
        assert r.status_code == 409 and "Replace was refused, so nothing changed" in r.json()["detail"]
        assert rows[0]["quote_number"] in r.json()["detail"]
        with SessionLocal() as db:
            assert db.query(CalculationRecord).filter_by(id=target).first().deleted_at is None
    finally:
        with SessionLocal() as db:
            db.execute(text("DELETE FROM icb_mes.production_jobs WHERE id = :j"), {"j": job})
            db.commit()


# ── 5. the results page, the documents ───────────────────────────────────────
def test_the_results_page_names_the_rev_withholds_generate_quote_and_says_sap(client, admin, chain, monkeypatch):
    from app import report_engine
    monkeypatch.setattr(report_engine, "_LETTERHEAD_OK", False)        # no WeasyPrint (today: :8000 and prod)
    html = client.get(f"/results/{chain['first']}", headers=admin).text
    assert f"{chain['qn']} · REV 0" in html and "REV 5 is the latest" in html
    assert "Official quotes are issued from SAP." in html
    assert f"/results/{chain['first']}/report" not in html and "Generate Quote" not in html
    latest = client.get(f"/results/{chain['r5']}", headers=admin).text
    assert f"{chain['qn']} · REV 5" in latest and 'data-testid="results-rev-latest"' in latest
    # the page's own Accept asks first on an older revision, and only there (the negative control: the latest)
    assert "This is REV 0. REV 5 is the latest. Accept REV 0 anyway?" in html and "Accept an older revision" in html
    assert "Accept an older revision" not in latest


def test_the_costing_exports_show_rev_from_rev_1_and_nothing_on_rev_0(client, admin, chain):
    from docx import Document
    from openpyxl import load_workbook

    def first_cell(rid):
        r = client.get(f"/results/{rid}/export/excel", headers=admin)
        assert r.status_code == 200
        ws = load_workbook(io.BytesIO(r.content)).worksheets[0]
        return next(str(c.value) for row in ws.iter_rows(max_row=6) for c in row if c.value not in (None, ""))

    def first_para(rid):
        r = client.get(f"/results/{rid}/export/word", headers=admin)
        assert r.status_code == 200
        return next(p.text for p in Document(io.BytesIO(r.content)).paragraphs if p.text.strip())
    assert first_cell(chain["first"]).startswith(f"{chain['qn']} — ") and "REV" not in first_cell(chain["first"])
    assert first_cell(chain["r1"]).startswith(f"{chain['qn']} · REV 1 — ")
    assert first_para(chain["r2"]).startswith(f"{chain['qn']} · REV 2 — ")
    assert "REV" not in first_para(chain["first"])


def test_the_letterhead_reference_carries_rev_from_rev_1(client, admin, ids, chain, monkeypatch):
    if not ids["has_template"]:
        pytest.skip("no report template in this database")
    from app import report_engine
    got = []
    monkeypatch.setattr(report_engine, "render_by_slug", lambda **kw: got.append(kw["quote_number"]) or b"%PDF-1.4")
    for rid in (chain["first"], chain["r1"]):
        assert client.get(f"/results/{rid}/report", headers=admin).status_code == 200
    assert got == [chain["qn"], f"{chain['qn']} · REV 1"]


def test_the_prejob_card_payloads_carry_the_revision_label():
    # 23 Oct: the label was first added to the wrong schema class, and every Pre-Job Card create 500'd in CI
    from app.schemas.prejob import OutstandingSignoffOut, PrejobCardOut
    assert "quote_label" in PrejobCardOut.model_fields and "quote_label" in OutstandingSignoffOut.model_fields
    out = PrejobCardOut.model_construct()
    out.quote_label = "A1/10/2026 · REV 1"                # what routers/prejob_cards._out sets
    assert out.quote_label.endswith("REV 1")



# ── 6. RT10_RULING_2 — "Generate Quote" is capability-gated (Michael chose option B; RT11 installs WeasyPrint) ──────
def _gate_problems(html: str, available: bool, record_id: int) -> list[str]:
    """What is wrong with the results page's quote area for a body costing, given the flag."""
    probs = []
    button = 'data-testid="generate-quote-btn"' in html
    link = f"/results/{record_id}/report" in html
    sap = "Official quotes are issued from SAP." in html
    if available:
        if not (button and link):
            probs.append("the Generate Quote button is missing while the letterhead can render")
        if sap:
            probs.append("the SAP line shows while the letterhead can render")
    else:
        if button or link:
            probs.append("the Generate Quote button shows while the letterhead cannot render")
        if not sap:
            probs.append("the SAP line is missing while the letterhead cannot render")
    return probs


@pytest.mark.parametrize("available", [True, False])
def test_the_results_page_offers_generate_quote_only_where_the_letterhead_can_render(client, admin, ids, chain,
                                                                                   monkeypatch, available):
    if not ids["has_template"]:
        pytest.skip("no report template in this database")
    from app import report_engine
    monkeypatch.setattr(report_engine, "_LETTERHEAD_OK", available)
    html = client.get(f"/results/{chain['first']}", headers=admin).text
    assert _gate_problems(html, available, chain["first"]) == []


def test_negative_control_a_button_shown_without_the_capability_is_caught():
    # the page as it would be if the template ignored the flag: the button AND no SAP line, with the flag False
    shown = '<a href="/results/7/report" data-testid="generate-quote-btn">📑 Generate Quote</a>'
    assert _gate_problems(shown, False, 7) == [
        "the Generate Quote button shows while the letterhead cannot render",
        "the SAP line is missing while the letterhead cannot render"]


class _FakeHTML:
    result = b"%PDF-1.7\n%fake"
    raises = False

    def __init__(self, string=None, **kw):
        self.string = string

    def write_pdf(self):
        if _FakeHTML.raises:
            raise OSError("cannot load library 'libpango-1.0-0'")
        return _FakeHTML.result


@pytest.mark.parametrize("case, want", [("renders", True), ("raises", False), ("not_a_pdf", False), ("missing", False)])
def test_the_capability_is_weasyprint_plus_a_one_page_render(monkeypatch, case, want):
    import sys
    import types
    from app import report_engine
    fake = types.ModuleType("weasyprint")
    fake.HTML = _FakeHTML
    monkeypatch.setattr(_FakeHTML, "raises", case == "raises")
    monkeypatch.setattr(_FakeHTML, "result", b"<html>not a pdf" if case == "not_a_pdf" else b"%PDF-1.7\n%fake")
    monkeypatch.setitem(sys.modules, "weasyprint", None if case == "missing" else fake)   # None = ImportError
    monkeypatch.setattr(report_engine, "_LETTERHEAD_OK", None)
    assert report_engine.check_letterhead_quote() is want
    assert report_engine.letterhead_quote_available() is want          # cached: no second render
